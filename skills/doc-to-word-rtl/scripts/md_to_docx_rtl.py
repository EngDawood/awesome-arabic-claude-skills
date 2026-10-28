#!/usr/bin/env python3
"""
Claude Doc markdown export -> house-style RTL .docx.

Usage: python3 md_to_docx_rtl.py INPUT OUTPUT.docx [options]
INPUT: a .md file, or the JSON file Claude Docs `export` saves ({"data": {"bytes_b64": ...}}).
"""
import argparse, base64, json, os, re, sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# Arabic Unicode blocks, built from code points so no invisible characters live in this file.
_AR_RANGES = [(0x0600, 0x06FF), (0x0750, 0x077F), (0x08A0, 0x08FF), (0xFB50, 0xFDFF), (0xFE70, 0xFEFF)]
AR_RE = re.compile('[' + ''.join(re.escape(chr(a)) + '-' + re.escape(chr(b)) for a, b in _AR_RANGES) + ']')
LATIN_RE = re.compile(r'[A-Za-z]')
BYLINE_RE = re.compile(r'^\s*([A-Z][a-z]{2} \d{1,2}, \d{4}|\d{4}-\d{2}-\d{2})?\s*(·\s*)?(@\S+\s*)*$')
INLINE_RE = re.compile(r'(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')
CENTER_H_RE = re.compile(r'^(المحور|الفصل|مراجع|المراجع|قائمة|ملخص|Chapter|References|Abstract)')
CAPTION_RE = re.compile(r'^(جدول|شكل|الجدول|الشكل|Table|Figure)\s*\(?\s*\d')
# Start-aligned = no jc at all. In a right-to-left paragraph Word reads jc="right" as the line's
# end, i.e. the LEFT edge; leaving jc unset aligns to the start, which is the right edge.
START = None
SOURCE_RE = re.compile(r'^(المصدر|Source)\s*:')


def load_markdown(path):
    raw = open(path, encoding='utf-8').read()
    s = raw.lstrip()
    if s.startswith('{'):
        try:
            d = json.loads(s)
            data = d.get('data', d)
            if 'bytes_b64' in data:
                return base64.b64decode(data['bytes_b64']).decode('utf-8')
            if 'markdown' in data:
                return data['markdown']
        except Exception:
            pass
    return raw


def is_rtl_text(t, default_rtl):
    ar = len(AR_RE.findall(t))
    la = len(LATIN_RE.findall(t))
    if ar == 0 and la == 0:
        return default_rtl
    return ar >= la * 0.3 if ar else False


def set_par_dir(p, rtl):
    pPr = p._p.get_or_add_pPr()
    for el in pPr.findall(qn('w:bidi')):
        pPr.remove(el)
    if rtl:
        b = OxmlElement('w:bidi'); b.set(qn('w:val'), '1')
        pPr.insert(1 if pPr.find(qn('w:pStyle')) is not None else 0, b)


def style_run(run, cfg, rtl, size=None, bold=None, italic=None, mono=False):
    rPr = run._r.get_or_add_rPr()
    fonts = rPr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rPr.insert(0, fonts)
    en = 'Consolas' if mono else cfg.en_font
    fonts.set(qn('w:ascii'), en); fonts.set(qn('w:hAnsi'), en)
    fonts.set(qn('w:cs'), cfg.ar_font); fonts.set(qn('w:eastAsia'), en)
    sz = size or cfg.size
    run.font.size = Pt(sz)
    szCs = rPr.find(qn('w:szCs'))
    if szCs is None:
        szCs = OxmlElement('w:szCs'); rPr.append(szCs)
    szCs.set(qn('w:val'), str(int(sz * 2)))
    if bold is not None:
        run.bold = bold
        if bold:
            rPr.append(OxmlElement('w:bCs'))
    if italic is not None:
        run.italic = italic
        if italic:
            rPr.append(OxmlElement('w:iCs'))
    if rtl and AR_RE.search(run.text or ''):
        r = OxmlElement('w:rtl'); r.set(qn('w:val'), '1'); rPr.append(r)
    run.font.color.rgb = RGBColor.from_string(cfg.color)


def add_inline(p, text, cfg, rtl, size=None, bold=False):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        b, i, mono, t = bold, False, False, part
        if (part.startswith('**') and part.endswith('**')) or (part.startswith('__') and part.endswith('__')):
            b, t = True, part[2:-2]
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            i, t = True, part[1:-1]
        elif part.startswith('`') and part.endswith('`'):
            mono, t = True, part[1:-1]
        elif part.startswith('[') and '](' in part:
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part)
            t = m.group(1) if m else part
        t = t.replace('\\*', '*').replace('\\_', '_').replace('\\[', '[').replace('\\]', ']').replace('\\#', '#').replace('\\-', '-').replace('\\.', '.')
        run = p.add_run(t)
        style_run(run, cfg, rtl, size=size, bold=b, italic=i, mono=mono)


def para_spacing(p, cfg, before=0, after=8, line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.line_spacing = line if line is not None else cfg.line


def cell_shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def table_borders(table):
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '8'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'auto')
        borders.append(e)
    table._tbl.tblPr.append(borders)


def table_rtl(table):
    bv = OxmlElement('w:bidiVisual'); bv.set(qn('w:val'), '1')
    table._tbl.tblPr.insert(0, bv)


def no_split(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    if header:
        h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
    cs = OxmlElement('w:cantSplit'); cs.set(qn('w:val'), 'true'); trPr.append(cs)


def add_page_number(section, cfg):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if kind:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); run._r.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; run._r.append(it)
    style_run(run, cfg, False, size=11)


def setup_document(cfg):
    doc = Document()
    sec = doc.sections[0]
    if cfg.paper == 'a4':
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    else:
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.orientation = WD_ORIENT.PORTRAIT
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Cm(cfg.margins))
    if cfg.doc_rtl:
        sec._sectPr.append(OxmlElement('w:bidi'))
    st = doc.styles['Normal']
    st.font.name = cfg.en_font; st.font.size = Pt(cfg.size)
    rPr = st.element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rPr.insert(0, rf)
    for a in ('ascii', 'hAnsi', 'eastAsia'):
        rf.set(qn(f'w:{a}'), cfg.en_font)
    rf.set(qn('w:cs'), cfg.ar_font)
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'en-US'); lang.set(qn('w:bidi'), 'ar-SA'); rPr.append(lang)
    szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), str(int(cfg.size * 2))); rPr.append(szcs)
    for lvl, sz in ((1, cfg.size + 6), (2, cfg.size + 2), (3, cfg.size), (4, cfg.size)):
        hs = doc.styles[f'Heading {lvl}']
        hs.font.size = Pt(sz); hs.font.bold = True; hs.font.color.rgb = RGBColor.from_string(cfg.color)
        hs.font.name = cfg.en_font
        hr = hs.element.get_or_add_rPr()
        f = hr.find(qn('w:rFonts'))
        if f is None:
            f = OxmlElement('w:rFonts'); hr.insert(0, f)
        for a in ('ascii', 'hAnsi', 'eastAsia'):
            f.set(qn(f'w:{a}'), cfg.en_font)
        f.set(qn('w:cs'), cfg.ar_font)
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if f.get(qn(a)) is not None:
                del f.attrib[qn(a)]
    add_page_number(sec, cfg)
    return doc


def add_heading(doc, text, level, cfg):
    rtl = is_rtl_text(text, cfg.doc_rtl)
    p = doc.add_paragraph(style=f'Heading {min(level, 4)}')
    set_par_dir(p, rtl)
    sizes = {1: cfg.size + 6, 2: cfg.size + 2, 3: cfg.size, 4: cfg.size}
    add_inline(p, text, cfg, rtl, size=sizes[min(level, 4)], bold=True)
    centered = level == 1 or bool(CENTER_H_RE.match(text))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else START
    para_spacing(p, cfg, before=10 if level <= 2 else 6, after=6)
    p.paragraph_format.keep_with_next = True


def add_paragraph(doc, text, cfg):
    rtl = is_rtl_text(text, cfg.doc_rtl)
    p = doc.add_paragraph()
    set_par_dir(p, rtl)
    if CAPTION_RE.match(text):
        add_inline(p, text, cfg, rtl, bold=True, size=cfg.size)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para_spacing(p, cfg, before=6, after=4)
        p.paragraph_format.keep_with_next = True
        return
    if SOURCE_RE.match(text):
        add_inline(p, text, cfg, rtl, size=cfg.size - 2)
        p.alignment = START
        para_spacing(p, cfg, before=4, after=10)
        return
    add_inline(p, text, cfg, rtl, size=cfg.size)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para_spacing(p, cfg, before=10 if cfg.after_table else 0)


def add_reference(doc, text, cfg):
    rtl = is_rtl_text(text, cfg.doc_rtl)
    p = doc.add_paragraph()
    set_par_dir(p, rtl)
    add_inline(p, text, cfg, rtl, size=cfg.size)
    p.alignment = START
    pf = p.paragraph_format
    pf.left_indent = Cm(0.635); pf.first_line_indent = Cm(-0.635)
    para_spacing(p, cfg, after=6)


def new_num_restart(doc):
    """Fresh w:num for 'List Number' so each numbered list restarts at 1."""
    numbering = doc.part.numbering_part.element
    sid = doc.styles['List Number'].element.find('.//' + qn('w:numId')).get(qn('w:val'))
    absid = None
    for num in numbering.findall(qn('w:num')):
        if num.get(qn('w:numId')) == sid:
            absid = num.find(qn('w:abstractNumId')).get(qn('w:val'))
    new_id = max(int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))) + 1
    num = OxmlElement('w:num'); num.set(qn('w:numId'), str(new_id))
    a = OxmlElement('w:abstractNumId'); a.set(qn('w:val'), absid); num.append(a)
    ov = OxmlElement('w:lvlOverride'); ov.set(qn('w:ilvl'), '0')
    so = OxmlElement('w:startOverride'); so.set(qn('w:val'), '1'); ov.append(so); num.append(ov)
    numbering.append(num)
    return new_id


def add_list_item(doc, text, ordered, level, cfg, number):
    rtl = is_rtl_text(text, cfg.doc_rtl)
    style = 'List Number' if ordered else 'List Bullet'
    if level > 0:
        style += f' {min(level + 1, 3)}'
    try:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.left_indent = None
    except KeyError:
        p = doc.add_paragraph()
        text = (f'{number}. ' if ordered else '• ') + text
    if ordered and level == 0 and p.style.name.startswith('List Number'):
        if number == 1 or getattr(cfg, '_cur_num', None) is None:
            cfg._cur_num = new_num_restart(doc)
        pPr = p._p.get_or_add_pPr()
        numPr = OxmlElement('w:numPr')
        il = OxmlElement('w:ilvl'); il.set(qn('w:val'), '0'); numPr.append(il)
        ni = OxmlElement('w:numId'); ni.set(qn('w:val'), str(cfg._cur_num)); numPr.append(ni)
        pPr.insert(1, numPr)
    set_par_dir(p, rtl)
    add_inline(p, text, cfg, rtl, size=cfg.size)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para_spacing(p, cfg, after=4, line=max(cfg.line, 1.16))


# ---------------- images ----------------
# Claude Docs' markdown export drops pictures ("&#91;image: <alt>\]") and charts/widgets
# ("&#91;embedded content: <caption>\]"), leaving those placeholders.
# The skill fetches each picture and passes them in document order with --images (a JSON list of
# {"alt", "path"}; a chart's "alt" is its caption); plain markdown images "![alt](local/path.png)" also work.
IMG_RE = re.compile(r'^(?:&#91;|\\?\[)image:\s*(.*?)\\?\]$|^!\[([^\]]*)\]\(([^)\s]+)\)$'
                    r'|^(?:&#91;|\\?\[)embedded content:\s*(.*?)\\?\]$')


def load_images(path):
    if not path:
        return []
    data = json.load(open(path, encoding='utf-8'))
    return [d if isinstance(d, dict) else {'alt': '', 'path': d} for d in data]


def add_image(doc, cfg, src):
    """Centered picture, scaled down to the text width (and at most 18 cm tall). False if the file is missing."""
    from docx.image.image import Image as DocxImage
    if not src or not os.path.exists(src):
        return False
    img = DocxImage.from_file(src)
    w_cm = img.px_width / (img.horz_dpi or 96) * 2.54
    h_cm = img.px_height / (img.vert_dpi or 96) * 2.54
    usable = (21.0 if cfg.paper == 'a4' else 21.59) - 2 * cfg.margins
    scale = min(1.0, usable / w_cm, 18.0 / h_cm)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(6); pf.space_after = Pt(4); pf.line_spacing = 1.0
    pf.keep_with_next = True
    p.add_run().add_picture(src, width=Cm(w_cm * scale))
    return True


def handle_image(doc, cfg, match, placeholder):
    """Returns (placeholder_text or None, caption or None) for the caller to write as paragraphs."""
    caption = match.group(4)
    alt = (match.group(1) or match.group(2) or caption or '').strip()
    src = match.group(3)
    if src is None:
        entry = cfg.images.pop(0) if cfg.images else None
        src = entry.get('path') if entry else None
        if entry and entry.get('alt') and alt and entry['alt'].strip() != alt:
            print(f'warning: image order mismatch: placeholder "{alt[:40]}" got "{entry["alt"][:40]}"', file=sys.stderr)
    if add_image(doc, cfg, src):
        cfg.img_ok += 1
        return None, (caption.strip() if caption else None)
    cfg.img_missing.append(alt[:40])
    return placeholder.format(alt=alt), None


def split_row(line):
    s = line.strip()
    if s.startswith('|'): s = s[1:]
    if s.endswith('|'): s = s[:-1]
    cells, cur, esc = [], '', False
    for ch in s:
        if esc:
            cur += ch; esc = False
        elif ch == '\\':
            esc = True
        elif ch == '|':
            cells.append(cur.strip()); cur = ''
        else:
            cur += ch
    cells.append(cur.strip())
    return cells


def add_table(doc, rows, cfg):
    header, body = rows[0], rows[1:]
    ncols = max(len(r) for r in rows)
    rtl = is_rtl_text(' '.join(' '.join(r) for r in rows), cfg.doc_rtl)
    t = doc.add_table(rows=1 + len(body), cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(t)
    if rtl:
        table_rtl(t)
    for ri, r in enumerate([header] + body):
        row = t.rows[ri]
        no_split(row, header=(ri == 0))
        for ci in range(ncols):
            txt = r[ci] if ci < len(r) else ''
            cell = row.cells[ci]
            p = cell.paragraphs[0]
            crtl = is_rtl_text(txt, rtl)
            set_par_dir(p, crtl)
            add_inline(p, txt, cfg, crtl, size=cfg.table_size, bold=(ri == 0))
            numeric = bool(re.fullmatch(r'[\d٠-٩.,%\s/\-–]+[مهـ]*', txt)) or txt == ''
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (ri == 0 or numeric) else START
            para_spacing(p, cfg, before=2, after=2, line=1.0)
            tcPr = cell._tc.get_or_add_tcPr()
            va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center'); tcPr.append(va)
            if ri == 0:
                cell_shade(cell, 'E7E6E6')
    usable = (21.0 if cfg.paper == 'a4' else 21.59) - 2 * cfg.margins
    lens = [max(len(r[c]) if c < len(r) else 0 for r in rows) for c in range(ncols)]
    weights = [min(max(l, 8), 50) for l in lens]
    tot = sum(weights)
    widths = [max(w / tot, 0.18) for w in weights]
    s = sum(widths); widths = [w / s * usable for w in widths]
    t.autofit = False
    for ci, gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths[ci] / 2.54 * 1440)))
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); t._tbl.tblPr.append(lay)
    for row in t.rows:
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
    mar = OxmlElement('w:tblCellMar')
    for side in ('left', 'right'):
        e = OxmlElement(f'w:{side}'); e.set(qn('w:w'), '100'); e.set(qn('w:type'), 'dxa'); mar.append(e)
    t._tbl.tblPr.append(mar)
    cfg.after_table = True


def convert(md, out, cfg):
    doc = setup_document(cfg)
    lines = md.replace('\r\n', '\n').split('\n')
    i, n = 0, len(lines)
    in_refs = False
    ol_counter = 0
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1; ol_counter = 0; continue
        if s in ('---', '***', '___'):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1; continue
        if s.startswith('```'):
            i += 1; buf = []
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph(); set_par_dir(p, False)
            r = p.add_run('\n'.join(buf)); style_run(r, cfg, False, size=10, mono=True)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            continue
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            cfg.after_table = False
            level = len(m.group(1)); text = m.group(2).strip()
            add_heading(doc, text, level, cfg)
            in_refs = bool(re.search(r'(المراجع|مراجع|References|Bibliography|المصادر)', text))
            if level == 1 and not cfg.keep_byline:
                j = i + 1
                while j < n and not lines[j].strip():
                    j += 1
                if j < n and ((BYLINE_RE.match(lines[j].strip()) and '·' in lines[j]) or re.match(r'^\s*@\S+\s*$', lines[j])):
                    i = j
            i += 1; continue
        im = IMG_RE.match(s)
        if im:
            cfg.after_table = False
            text, caption = handle_image(doc, cfg, im, '[صورة: {alt}]')
            for t in (text, caption):
                if t:
                    add_paragraph(doc, t, cfg)
            i += 1; continue
        if s.startswith('|') and i + 1 < n and re.match(r'^\s*\|?\s*:?-{2,}', lines[i + 1]):
            rows = [split_row(s)]
            i += 2
            while i < n and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i])); i += 1
            add_table(doc, rows, cfg)
            continue
        lm = re.match(r'^(\s*)([-*+]|\d+[.)])\s+(.*)$', line)
        if lm:
            indent = len(lm.group(1).replace('\t', '    ')) // 2
            ordered = lm.group(2)[0].isdigit()
            ol_counter = ol_counter + 1 if ordered else ol_counter
            add_list_item(doc, lm.group(3).strip(), ordered, min(indent, 2), cfg, ol_counter)
            i += 1; continue
        if s.startswith('>'):
            add_paragraph(doc, s.lstrip('>').strip(), cfg)
            i += 1; continue
        buf = [s]; i += 1
        while i < n and lines[i].strip() and not re.match(r'^(#{1,6}\s|\||[-*+]\s|\d+[.)]\s|```|>)', lines[i].strip()):
            buf.append(lines[i].strip()); i += 1
        text = ' '.join(buf)
        if in_refs and len(text) > 40:
            add_reference(doc, text, cfg)
        else:
            add_paragraph(doc, text, cfg)
        if not SOURCE_RE.match(text):
            cfg.after_table = False
    doc.core_properties.title = next((l[2:].strip() for l in lines if l.startswith('# ')), '')
    doc.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('input'); ap.add_argument('output')
    ap.add_argument('--ar-font', default='Traditional Arabic')
    ap.add_argument('--en-font', default='Traditional Arabic')
    ap.add_argument('--size', type=float, default=16)
    ap.add_argument('--table-size', type=float, default=14)
    ap.add_argument('--color', default='0D0D0D')
    ap.add_argument('--line', type=float, default=1.08)
    ap.add_argument('--margins', type=float, default=2.54)
    ap.add_argument('--paper', default='a4', choices=['a4', 'letter'])
    ap.add_argument('--keep-byline', action='store_true')
    ap.add_argument('--dir', default='auto', choices=['auto', 'rtl', 'ltr'])
    ap.add_argument('--images', help='JSON list of {"alt", "path"} in document order')
    cfg = ap.parse_args()
    cfg.images = load_images(cfg.images)
    cfg.img_ok, cfg.img_missing = 0, []
    cfg.after_table = False
    md = load_markdown(cfg.input)
    if cfg.dir == 'auto':
        cfg.doc_rtl = len(AR_RE.findall(md)) > len(LATIN_RE.findall(md))
    else:
        cfg.doc_rtl = cfg.dir == 'rtl'
    convert(md, cfg.output, cfg)
    print(f'OK -> {cfg.output}  (direction: {"RTL" if cfg.doc_rtl else "LTR"})')
    print(f'images: {cfg.img_ok} embedded, {len(cfg.img_missing)} missing'
          + (f' ({", ".join(cfg.img_missing)})' if cfg.img_missing else ''))


if __name__ == '__main__':
    main()
