#!/usr/bin/env python3
"""
Claude Doc markdown export -> English .docx in APA 7 student-paper style.

Usage: python3 md_to_docx_en.py INPUT OUTPUT.docx [options]
INPUT: a .md file, or the JSON file Claude Docs `export` saves ({"data": {"bytes_b64": ...}}).

Heading map: #  -> paper title (bold, centered)
             ## -> APA Level 1 (bold, centered)
             ### -> Level 2 (bold, flush left)
             #### -> Level 3 (bold italic, flush left)
             ##### -> Level 4 (bold, indented)
"""
import argparse, base64, json, os, re, sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BYLINE_RE = re.compile(r'^\s*([A-Z][a-z]{2} \d{1,2}, \d{4}|\d{4}-\d{2}-\d{2})?\s*(·\s*)?(@\S+\s*)*$')
INLINE_RE = re.compile(r'(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')
CAPTION_RE = re.compile(r'^\**(Table|Figure)\s+(\d+)\**\s*[:.\-–]?\s*(.*)$')
NOTE_RE = re.compile(r'^\**(Note|Source)\.?\**\s*[.:]\s*(.*)$')
REFS_RE = re.compile(r'^(References|Reference List|Bibliography|Works Cited)\b', re.I)
INDENT = Cm(1.27)  # 0.5 inch


def load_markdown(path):
    raw = open(path, encoding='utf-8').read()
    s = raw.lstrip()
    if s.startswith('{'):
        try:
            d = json.loads(s)
            data = d.get('data', d)
            if 'bytes_b64' in data:
                return base64.b64decode(data['bytes_b64']).decode('utf-8')
            if 'markdown' in data:
                return data['markdown']
        except Exception:
            pass
    return raw


def style_run(run, cfg, size=None, bold=None, italic=None, mono=False):
    rPr = run._r.get_or_add_rPr()
    fonts = rPr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rPr.insert(0, fonts)
    f = 'Consolas' if mono else cfg.font
    for a in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn(f'w:{a}'), f)
    run.font.size = Pt(size or cfg.size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_inline(p, text, cfg, size=None, bold=False, italic=False):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        b, i, mono, t = bold, italic, False, part
        if (part.startswith('**') and part.endswith('**')) or (part.startswith('__') and part.endswith('__')):
            b, t = True, part[2:-2]
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            i, t = (not italic), part[1:-1]
        elif part.startswith('`') and part.endswith('`'):
            mono, t = True, part[1:-1]
        elif part.startswith('[') and '](' in part:
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', part)
            t = m.group(1) if m else part
        t = t.replace('\\*', '*').replace('\\_', '_').replace('\\[', '[').replace('\\]', ']').replace('\\#', '#').replace('\\-', '-').replace('\\.', '.')
        style_run(p.add_run(t), cfg, size=size, bold=b, italic=i, mono=mono)


def fmt(p, cfg, align=WD_ALIGN_PARAGRAPH.LEFT, first=None, left=None, line=None, before=0, after=0):
    pf = p.paragraph_format
    p.alignment = align
    pf.space_before = Pt(before); pf.space_after = Pt(after)
    pf.line_spacing = line if line is not None else cfg.line
    if first is not None:
        pf.first_line_indent = first
    if left is not None:
        pf.left_indent = left


def border(el_parent, tag, edges):
    b = OxmlElement(tag)
    for edge, val in edges:
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), val); e.set(qn('w:sz'), '8'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'auto')
        b.append(e)
    el_parent.append(b)


def add_page_number(section, cfg):
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    for kind, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if kind:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); run._r.append(fc)
        else:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; run._r.append(it)
    style_run(run, cfg)


def setup_document(cfg):
    doc = Document()
    sec = doc.sections[0]
    if cfg.paper == 'a4':
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    else:
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.orientation = WD_ORIENT.PORTRAIT
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Cm(cfg.margins))
    st = doc.styles['Normal']
    st.font.name = cfg.font; st.font.size = Pt(cfg.size)
    rf = st.element.get_or_add_rPr().find(qn('w:rFonts'))
    for a in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn(f'w:{a}'), cfg.font)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = cfg.line
    for lvl in range(1, 6):
        hs = doc.styles[f'Heading {lvl}']
        hs.font.size = Pt(cfg.size); hs.font.color.rgb = RGBColor(0, 0, 0); hs.font.name = cfg.font
        f = hs.element.get_or_add_rPr().find(qn('w:rFonts'))
        if f is None:
            f = OxmlElement('w:rFonts'); hs.element.get_or_add_rPr().insert(0, f)
        for a in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            f.set(qn(f'w:{a}'), cfg.font)
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            if f.get(qn(a)) is not None:
                del f.attrib[qn(a)]
    add_page_number(sec, cfg)
    return doc


def add_heading(doc, text, md_level, cfg):
    # md '#' = title; '##'..'#####' = APA levels 1..4
    apa = md_level - 1
    p = doc.add_paragraph(style=f'Heading {min(max(apa, 1), 5)}')
    if apa <= 0:
        add_inline(p, text, cfg, bold=True); fmt(p, cfg, WD_ALIGN_PARAGRAPH.CENTER)
    elif apa == 1:
        add_inline(p, text, cfg, bold=True); fmt(p, cfg, WD_ALIGN_PARAGRAPH.CENTER)
    elif apa == 2:
        add_inline(p, text, cfg, bold=True); fmt(p, cfg)
    elif apa == 3:
        add_inline(p, text, cfg, bold=True, italic=True); fmt(p, cfg)
    else:
        add_inline(p, text, cfg, bold=True); fmt(p, cfg, first=INDENT)
    p.paragraph_format.keep_with_next = True


def add_body(doc, text, cfg):
    p = doc.add_paragraph()
    add_inline(p, text, cfg)
    fmt(p, cfg, first=INDENT)


def add_caption(doc, kind, num, title, cfg):
    p = doc.add_paragraph(); add_inline(p, f'{kind} {num}', cfg, bold=True); fmt(p, cfg)
    p.paragraph_format.keep_with_next = True
    if title:
        p = doc.add_paragraph(); add_inline(p, title.strip('*'), cfg, italic=True); fmt(p, cfg)
        p.paragraph_format.keep_with_next = True


def add_note(doc, label, text, cfg):
    p = doc.add_paragraph()
    style_run(p.add_run(f'{label}.'), cfg, italic=True)
    add_inline(p, ' ' + text, cfg)
    fmt(p, cfg, after=cfg.size)


def add_reference(doc, text, cfg):
    p = doc.add_paragraph()
    add_inline(p, text, cfg)
    fmt(p, cfg, left=INDENT, first=-INDENT)


def new_num_restart(doc):
    numbering = doc.part.numbering_part.element
    sid = doc.styles['List Number'].element.find('.//' + qn('w:numId')).get(qn('w:val'))
    absid = None
    for num in numbering.findall(qn('w:num')):
        if num.get(qn('w:numId')) == sid:
            absid = num.find(qn('w:abstractNumId')).get(qn('w:val'))
    new_id = max(int(n.get(qn('w:numId'))) for n in numbering.findall(qn('w:num'))) + 1
    num = OxmlElement('w:num'); num.set(qn('w:numId'), str(new_id))
    a = OxmlElement('w:abstractNumId'); a.set(qn('w:val'), absid); num.append(a)
    ov = OxmlElement('w:lvlOverride'); ov.set(qn('w:ilvl'), '0')
    so = OxmlElement('w:startOverride'); so.set(qn('w:val'), '1'); ov.append(so); num.append(ov)
    numbering.append(num)
    return new_id


def add_list_item(doc, text, ordered, level, cfg, number):
    style = 'List Number' if ordered else 'List Bullet'
    if level > 0:
        style += f' {min(level + 1, 3)}'
    p = doc.add_paragraph(style=style)
    if ordered and level == 0:
        if number == 1 or getattr(cfg, '_cur_num', None) is None:
            cfg._cur_num = new_num_restart(doc)
        numPr = OxmlElement('w:numPr')
        il = OxmlElement('w:ilvl'); il.set(qn('w:val'), '0'); numPr.append(il)
        ni = OxmlElement('w:numId'); ni.set(qn('w:val'), str(cfg._cur_num)); numPr.append(ni)
        p._p.get_or_add_pPr().insert(1, numPr)
    add_inline(p, text, cfg)
    fmt(p, cfg)


# ---------------- images ----------------
# Claude Docs' markdown export drops pictures ("&#91;image: <alt>\]") and charts/widgets
# ("&#91;embedded content: <caption>\]"), leaving those placeholders.
# The skill fetches each picture and passes them in document order with --images (a JSON list of
# {"alt", "path"}; a chart's "alt" is its caption); plain markdown images "![alt](local/path.png)" also work.
IMG_RE = re.compile(r'^(?:&#91;|\\?\[)image:\s*(.*?)\\?\]$|^!\[([^\]]*)\]\(([^)\s]+)\)$'
                    r'|^(?:&#91;|\\?\[)embedded content:\s*(.*?)\\?\]$')


def load_images(path):
    if not path:
        return []
    data = json.load(open(path, encoding='utf-8'))
    return [d if isinstance(d, dict) else {'alt': '', 'path': d} for d in data]


def add_image(doc, cfg, src):
    """Centered picture, scaled down to the text width (and at most 18 cm tall). False if the file is missing."""
    from docx.image.image import Image as DocxImage
    if not src or not os.path.exists(src):
        return False
    img = DocxImage.from_file(src)
    w_cm = img.px_width / (img.horz_dpi or 96) * 2.54
    h_cm = img.px_height / (img.vert_dpi or 96) * 2.54
    usable = (21.0 if cfg.paper == 'a4' else 21.59) - 2 * cfg.margins
    scale = min(1.0, usable / w_cm, 18.0 / h_cm)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(6); pf.space_after = Pt(4); pf.line_spacing = 1.0
    pf.keep_with_next = True
    p.add_run().add_picture(src, width=Cm(w_cm * scale))
    return True


def handle_image(doc, cfg, match, placeholder):
    """Returns (placeholder_text or None, caption or None) for the caller to write as paragraphs."""
    caption = match.group(4)
    alt = (match.group(1) or match.group(2) or caption or '').strip()
    src = match.group(3)
    if src is None:
        entry = cfg.images.pop(0) if cfg.images else None
        src = entry.get('path') if entry else None
        if entry and entry.get('alt') and alt and entry['alt'].strip() != alt:
            print(f'warning: image order mismatch: placeholder "{alt[:40]}" got "{entry["alt"][:40]}"', file=sys.stderr)
    if add_image(doc, cfg, src):
        cfg.img_ok += 1
        return None, (caption.strip() if caption else None)
    cfg.img_missing.append(alt[:40])
    return placeholder.format(alt=alt), None


def split_row(line):
    s = line.strip()
    if s.startswith('|'): s = s[1:]
    if s.endswith('|'): s = s[:-1]
    cells, cur, esc = [], '', False
    for ch in s:
        if esc:
            cur += ch; esc = False
        elif ch == '\\':
            esc = True
        elif ch == '|':
            cells.append(cur.strip()); cur = ''
        else:
            cur += ch
    cells.append(cur.strip())
    return cells


def add_table(doc, rows, cfg):
    """APA table: horizontal rules only (top, under header, bottom), no shading."""
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    border(t._tbl.tblPr, 'w:tblBorders', [('top', 'single'), ('bottom', 'single'), ('left', 'nil'), ('right', 'nil'), ('insideH', 'nil'), ('insideV', 'nil')])
    for ri, r in enumerate(rows):
        row = t.rows[ri]
        trPr = row._tr.get_or_add_trPr()
        if ri == 0:
            h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
        cs = OxmlElement('w:cantSplit'); cs.set(qn('w:val'), 'true'); trPr.append(cs)
        for ci in range(ncols):
            txt = r[ci] if ci < len(r) else ''
            cell = row.cells[ci]
            p = cell.paragraphs[0]
            add_inline(p, txt, cfg, size=cfg.table_size)
            numeric = bool(re.fullmatch(r'[\d.,%\s/\-–()<>=*]+', txt)) or txt == ''
            fmt(p, cfg, WD_ALIGN_PARAGRAPH.CENTER if (ri == 0 or numeric) else WD_ALIGN_PARAGRAPH.LEFT, line=1.0, before=2, after=2)
            tcPr = cell._tc.get_or_add_tcPr()
            va = OxmlElement('w:vAlign'); va.set(qn('w:val'), 'center'); tcPr.append(va)
            if ri == 0:
                border(tcPr, 'w:tcBorders', [('bottom', 'single')])
    usable = (21.0 if cfg.paper == 'a4' else 21.59) - 2 * cfg.margins
    lens = [max(len(r[c]) if c < len(r) else 0 for r in rows) for c in range(ncols)]
    weights = [min(max(l, 8), 50) for l in lens]
    widths = [max(w / sum(weights), 0.15) for w in weights]
    s = sum(widths); widths = [w / s * usable for w in widths]
    t.autofit = False
    for ci, gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths[ci] / 2.54 * 1440)))
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); t._tbl.tblPr.append(lay)
    for row in t.rows:
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])


def convert(md, out, cfg):
    doc = setup_document(cfg)
    lines = md.replace('\r\n', '\n').split('\n')
    i, n = 0, len(lines)
    in_refs = False
    ol_counter = 0
    after_table = False
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1; ol_counter = 0; continue
        if s in ('---', '***', '___'):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1; continue
        if s.startswith('```'):
            i += 1; buf = []
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i]); i += 1
            i += 1
            p = doc.add_paragraph(); style_run(p.add_run('\n'.join(buf)), cfg, size=10, mono=True); fmt(p, cfg, line=1.0)
            continue
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            after_table = False
            level = len(m.group(1)); text = m.group(2).strip()
            in_refs = bool(REFS_RE.match(text))
            if in_refs:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            add_heading(doc, text, level, cfg)
            if level == 1 and not cfg.keep_byline:
                j = i + 1
                while j < n and not lines[j].strip():
                    j += 1
                if j < n and ((BYLINE_RE.match(lines[j].strip()) and '·' in lines[j]) or re.match(r'^\s*@\S+\s*$', lines[j])):
                    i = j
            i += 1; continue
        im = IMG_RE.match(s)
        if im:
            after_table = False
            text, caption = handle_image(doc, cfg, im, '[Image: {alt}]')
            for t in (text, caption):
                if t:
                    add_body(doc, t, cfg)
            i += 1; continue
        if s.startswith('|') and i + 1 < n and re.match(r'^\s*\|?\s*:?-{2,}', lines[i + 1]):
            rows = [split_row(s)]
            i += 2
            while i < n and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i])); i += 1
            add_table(doc, rows, cfg)
            after_table = True
            continue
        lm = re.match(r'^(\s*)([-*+]|\d+[.)])\s+(.*)$', line)
        if lm:
            indent = len(lm.group(1).replace('\t', '    ')) // 2
            ordered = lm.group(2)[0].isdigit()
            ol_counter = ol_counter + 1 if ordered else ol_counter
            add_list_item(doc, lm.group(3).strip(), ordered, min(indent, 2), cfg, ol_counter)
            i += 1; continue
        buf = [s.lstrip('>').strip()]; i += 1
        while i < n and lines[i].strip() and not re.match(r'^(#{1,6}\s|\||[-*+]\s|\d+[.)]\s|```)', lines[i].strip()):
            buf.append(lines[i].strip()); i += 1
        text = ' '.join(buf)
        cm, nm = CAPTION_RE.match(text), NOTE_RE.match(text)
        if cm:
            add_caption(doc, cm.group(1), cm.group(2), cm.group(3), cfg)
        elif nm and after_table:
            add_note(doc, nm.group(1), nm.group(2), cfg)
        elif in_refs:
            add_reference(doc, text, cfg)
        else:
            if after_table:
                doc.add_paragraph()  # blank double-spaced line after a table
            add_body(doc, text, cfg)
        after_table = False
    doc.core_properties.title = next((l[2:].strip() for l in lines if l.startswith('# ')), '')
    doc.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('input'); ap.add_argument('output')
    ap.add_argument('--font', default='Times New Roman')
    ap.add_argument('--size', type=float, default=12)
    ap.add_argument('--table-size', type=float, default=12)
    ap.add_argument('--line', type=float, default=2.0)
    ap.add_argument('--margins', type=float, default=2.54)
    ap.add_argument('--paper', default='letter', choices=['a4', 'letter'])
    ap.add_argument('--keep-byline', action='store_true')
    ap.add_argument('--images', help='JSON list of {"alt", "path"} in document order')
    cfg = ap.parse_args()
    cfg.images = load_images(cfg.images)
    cfg.img_ok, cfg.img_missing = 0, []
    convert(load_markdown(cfg.input), cfg.output, cfg)
    print(f'OK -> {cfg.output}  (style: APA 7, LTR)')
    print(f'images: {cfg.img_ok} embedded, {len(cfg.img_missing)} missing'
          + (f' ({", ".join(cfg.img_missing)})' if cfg.img_missing else ''))


if __name__ == '__main__':
    main()
